import os, pytest
from bs4 import BeautifulSoup

def fake_page(u):
    if "bad.example" in u: raise RuntimeError("timeout")
    return {"url": u, "title": "Acme " + u[-8:], "description": "Payroll software", "text": "Acme builds payroll software for small businesses. Pricing starts at 99 dollars per month gold plan.", "soup": BeautifulSoup("", "html.parser")}

def run(appm, prompt="Research 2 payroll tools", council="full"):
    mid = appm.new_mission(prompt, council=council); appm.run_mission(mid); return appm.get_m(mid)

def test_graph_critic_revise_loop_and_excel_verification(appm, llm, monkeypatch):
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: [{"title": "A", "url": "https://a.example/1111111", "snippet": ""}, {"title": "B", "url": "https://c.example/2222222", "snippet": ""}])
    monkeypatch.setattr(appm, "fetch_page", fake_page); llm.extractor = {"company": "Acme", "product": "payroll software small businesses", "pricing": "99 dollars per month gold plan", "target_customers": "Unknown"}
    llm.critic = [{"flags": [{"i": 0, "issue": "pricing unsupported"}]}, {"flags": []}]; m = run(appm)
    steps = [l["msg"] for l in m["log"] if l["step"] == "council"]
    assert any("revised 1 row" in s for s in steps) and sum("Critic" in s for s in steps) >= 2
    assert m["status"] in ("completed", "completed_with_warnings") and os.path.exists(os.path.join(appm.OUT, m["artifact"]))
    assert any("Verifier: re-read file" in l["msg"] for l in m["log"])

def test_phoenix_switches_source_when_a_page_fails(appm, monkeypatch):
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: [{"title": "B", "url": "https://bad.example/x", "snippet": ""}, {"title": "A", "url": "https://good.example/1111111", "snippet": ""}])
    monkeypatch.setattr(appm, "fetch_page", fake_page); monkeypatch.setattr(appm.time, "sleep", lambda s: None); m = run(appm, "Research 1 payroll tool")
    assert any(l["step"] == "recover" and "switching" in l["msg"] for l in m["log"]) and m["status"].startswith("completed")

def test_search_failure_is_reported_honestly(appm, monkeypatch):
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: []); m = run(appm)
    assert m["status"] == "failed" and "no results" in m["log"][-1]["msg"].lower()

def test_all_sources_unreachable_fails_cleanly(appm, monkeypatch):
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: [{"title": "B", "url": "https://bad.example/x", "snippet": ""}]); monkeypatch.setattr(appm, "fetch_page", fake_page); monkeypatch.setattr(appm.time, "sleep", lambda s: None)
    m = run(appm, "Research 1 thing"); assert m["status"] == "failed" and "No source pages" in m["log"][-1]["msg"]

def test_playwright_unavailable_falls_back_to_http_with_warning(appm, monkeypatch):
    monkeypatch.setattr(appm, "BROWSER_MODE", "playwright"); monkeypatch.setattr(appm, "launch_guarded", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no chromium")))
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: [{"title": "A", "url": "https://good.example/1111111", "snippet": ""}]); monkeypatch.setattr(appm, "fetch_page", fake_page)
    m = run(appm, "Research 1 payroll tool"); assert any(l["step"] == "browser" and l["status"] == "warn" for l in m["log"]) and m["status"].startswith("completed")

def test_unknown_fields_stay_unknown_not_guessed(appm, monkeypatch):
    monkeypatch.setattr(appm, "ddg", lambda q, n=8: [{"title": "A", "url": "https://good.example/1111111", "snippet": ""}]); monkeypatch.setattr(appm, "fetch_page", fake_page)
    m = run(appm, "Research 1 payroll tool"); assert m["status"] == "completed_with_warnings"
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(appm.OUT, m["artifact"]))["Comparison"]; row = [c.value for c in ws[2]]; assert "Unknown" in row

def test_excel_verifier_detects_tampering(appm, tmp_path):
    recs = [{"Company": "A", "Product": "p", "Pricing": "x", "Target Customers": "y", "Source URL": "https://a/", "Evidence": "e", "Status": "Complete", "Trust": "green", "Retrieved At": "t", "Note": ""}]
    p = str(tmp_path / "x.xlsx"); appm.write_xlsx(p, recs); assert appm.verify_xlsx(p, recs)[0] == []
    recs2 = [dict(recs[0], **{"Source URL": "https://other/"})]; assert appm.verify_xlsx(p, recs2)[0]

def test_pause_cancel_and_fork_with_assumptions(appm, client):
    from conftest import H_
    mid = appm.new_mission("Research 2 payroll tools"); assert appm.get_m(mid)["status"] == "awaiting_approval"
    f = client.post(f"/api/missions/{mid}/fork", headers=H_, json={"prompt": "Research 3 payroll tools", "assumptions": "India only"}).json
    assert appm.get_m(f["id"])["plan"]["assumptions"] == "India only" and appm.get_m(f["id"])["plan"]["forked_from"] == mid
    assert client.post(f"/api/missions/{mid}/cancel", headers=H_).status_code == 200 and appm.get_m(mid)["status"] == "cancelled"

def test_mission_input_validation(client):
    from conftest import H_
    assert client.post("/api/missions", headers=H_, json={"prompt": "hi"}).status_code == 400 and client.post("/api/agent", headers=H_, json={"goal": "x"}).status_code == 400
    assert client.post("/api/agent", headers=H_, json={"goal": "open the page and read it", "start_url": "http://127.0.0.1/"}).status_code == 400
